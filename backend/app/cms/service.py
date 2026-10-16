"""Knowledge workflow: who may do what, and what happens to the live index."""

import csv
import io
import re
from collections.abc import Awaitable, Callable
from dataclasses import dataclass
from typing import Any

from app.cms.store import OPEN_STATES, DocRecord, DocVersion, KnowledgeStore, now
from app.personas.prompts import INSTRUCTIONS_MAX_CHARS
from app.rag.documents import KnowledgeDoc
from app.staff.auth import Principal
from app.staff.store import StaffStore

PERSONA_VALUES = ("driver", "client", "customer", "all")
_ID_RE = re.compile(r"^[a-z0-9][a-z0-9-]{1,79}$")


class CMSError(Exception):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


@dataclass
class DraftInput:
    persona: str
    title: str
    body: str
    tags: list[str]
    visibility: str = "public"


class CMSService:
    audit_prefix = "kb"

    def __init__(self, store: KnowledgeStore, staff: StaffStore, on_publish: Callable[[], Awaitable[Any]],
                 four_eyes: bool = True):
        self.store = store
        self.staff = staff
        self.on_publish = on_publish  # re-index after the published set changes
        self.four_eyes = four_eyes

    async def _audit(self, p: Principal, action: str, doc_id: str, details: dict[str, Any] | None = None) -> None:
        await self.staff.audit(p.email, f"{self.audit_prefix}.{action}", doc_id, details)

    # ------------------------------------------------------------------ reads
    async def published_docs(self) -> list[KnowledgeDoc]:
        docs = {d.id: d for d in await self.store.list_docs()}
        out = []
        for doc_id, vs in (await self.store.latest_versions()).items():
            d = docs.get(doc_id)
            if not d or d.status != "active":
                continue
            out += [v.to_doc() for v in vs if v.state == "published"]
        return out

    async def overview(self) -> list[dict[str, Any]]:
        latest = await self.store.latest_versions()
        creators = await self.store.first_authors()
        rows = []
        for d in await self.store.list_docs():
            vs = latest.get(d.id, [])
            pub = next((v for v in vs if v.state == "published"), None)
            draft = next((v for v in vs if v.state in OPEN_STATES), None)
            head = draft or pub
            rows.append({
                "id": d.id, "status": d.status, "published_version": d.published_version,
                "persona": head.persona if head else None, "title": head.title if head else d.id,
                "visibility": head.visibility if head else None,
                "draft": {"version": draft.version, "state": draft.state, "author": draft.author,
                          "source": draft.source, "updated_at": draft.updated_at} if draft else None,
                "updated_at": (head.updated_at if head else None) or d.updated_at,
                # who added it / who touched it last, topics, when the live version was published
                "created_by": creators.get(d.id), "last_editor": head.author if head else None,
                "tags": head.tags if head else [],
                "published_at": (pub.reviewed_at or pub.updated_at) if pub else None,
            })
        return rows

    async def candidates(self) -> list["Candidate"]:
        """Published and in-progress documents (active only) for duplicate / contradiction checks."""
        from app.cms.quality import Candidate

        active = {d.id for d in await self.store.list_docs() if d.status == "active"}
        out = []
        for doc_id, vs in (await self.store.latest_versions()).items():
            if doc_id not in active:
                continue
            for v in vs:
                out.append(Candidate(v.doc_id, v.title, v.body, v.persona, v.state, v.author))
        return out

    async def review_queue(self) -> list[dict[str, Any]]:
        out = []
        for vs in (await self.store.latest_versions()).values():
            pub = next((v for v in vs if v.state == "published"), None)
            for v in vs:
                if v.state == "in_review":
                    out.append({"draft": v.to_dict(), "published": pub.to_dict() if pub else None})
        return sorted(out, key=lambda x: x["draft"]["updated_at"] or "")

    async def detail(self, doc_id: str) -> dict[str, Any]:
        d = await self._doc(doc_id)
        return {"doc": d.__dict__, "versions": [v.to_dict() for v in await self.store.versions(doc_id)]}

    # ----------------------------------------------------------------- writes
    async def save_draft(self, p: Principal, doc_id: str, data: DraftInput, source: str = "editor") -> DocVersion:
        self._validate(doc_id, data)
        d = await self.store.get_doc(doc_id)
        if not d:
            d = DocRecord(id=doc_id)
            await self.store.upsert_doc(d)
        vs = await self.store.versions(doc_id)
        open_v = next((v for v in vs if v.state in OPEN_STATES), None)
        ts = now()
        if open_v:
            open_v.persona, open_v.title, open_v.body = data.persona, data.title, data.body
            open_v.tags, open_v.visibility, open_v.updated_at = data.tags, data.visibility, ts
            # Editing a draft under review sends it back to draft: reviewers approve what they read.
            open_v.state, open_v.reviewer, open_v.reviewed_at = "draft", None, None
            await self.store.update_version(open_v)
            v = open_v
        else:
            v = DocVersion(doc_id=doc_id, version=(vs[0].version + 1) if vs else 1, persona=data.persona,
                           title=data.title, body=data.body, tags=data.tags, visibility=data.visibility,
                           state="draft", source=source, author=p.email, created_at=ts, updated_at=ts)
            await self.store.insert_version(v)
        await self._audit(p, "save_draft", doc_id, {"version": v.version, "source": source})
        return v

    async def submit(self, p: Principal, doc_id: str) -> DocVersion:
        v = await self._open(doc_id, ("draft", "rejected"))
        v.state, v.updated_at, v.review_comment = "in_review", now(), None
        await self.store.update_version(v)
        await self._audit(p, "submit", doc_id, {"version": v.version})
        return v

    async def approve(self, p: Principal, doc_id: str, version: int, comment: str | None) -> DocVersion:
        v = await self._version(doc_id, version)
        if v.state != "in_review":
            raise CMSError("only versions in review can be approved", 409)
        if self.four_eyes and v.author == p.email and "admin" not in p.roles:
            raise CMSError("a second person must approve your own change", 403)
        v.reviewer, v.reviewed_at, v.review_comment = p.email, now(), comment
        await self._publish(v)
        await self._audit(p, "approve", doc_id, {"version": version, "comment": comment})
        return v

    async def reject(self, p: Principal, doc_id: str, version: int, comment: str) -> DocVersion:
        v = await self._version(doc_id, version)
        if v.state != "in_review":
            raise CMSError("only versions in review can be rejected", 409)
        v.state, v.reviewer, v.reviewed_at, v.review_comment = "rejected", p.email, now(), comment
        await self.store.update_version(v)
        await self._audit(p, "reject", doc_id, {"version": version, "comment": comment})
        return v

    async def discard(self, p: Principal, doc_id: str) -> None:
        v = await self._open(doc_id, OPEN_STATES)
        await self.store.delete_version(doc_id, v.version)
        d = await self._doc(doc_id)
        if d.published_version is None and not await self.store.versions(doc_id):
            d.status = "archived"
            await self.store.upsert_doc(d)
        await self._audit(p, "discard_draft", doc_id, {"version": v.version})

    async def rollback(self, p: Principal, doc_id: str, to_version: int) -> DocVersion:
        """Re-publish an old version as a new version (history is never rewritten)."""
        old = await self._version(doc_id, to_version)
        vs = await self.store.versions(doc_id)
        if any(v.state in OPEN_STATES for v in vs):
            raise CMSError("discard or finish the open draft before rolling back", 409)
        ts = now()
        v = DocVersion(doc_id=doc_id, version=vs[0].version + 1, persona=old.persona, title=old.title,
                       body=old.body, tags=old.tags, visibility=old.visibility, state="draft", source="rollback",
                       author=p.email, created_at=ts, updated_at=ts, reviewer=p.email, reviewed_at=ts,
                       review_comment=f"rollback to v{to_version}")
        await self.store.insert_version(v)
        await self._publish(v)
        await self._audit(p, "rollback", doc_id, {"to_version": to_version, "new_version": v.version})
        return v

    async def set_archived(self, p: Principal, doc_id: str, archived: bool) -> None:
        d = await self._doc(doc_id)
        if not archived and d.published_version is None:
            raise CMSError("nothing published to restore", 409)
        d.status = "archived" if archived else "active"
        await self.store.upsert_doc(d)
        await self._audit(p, "archive" if archived else "restore", doc_id)
        await self.on_publish()

    async def import_rows(self, p: Principal, rows: list[dict[str, str]], submit: bool = True) -> dict[str, Any]:
        """Bulk FAQ import. Columns: id (optional), persona, title|question, body|answer, tags, visibility."""
        created, errors = [], []
        for i, row in enumerate(rows, start=2):  # row 1 = header
            r = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
            title = r.get("title") or r.get("question") or ""
            body = r.get("body") or r.get("answer") or ""
            persona = r.get("persona", "").lower()
            doc_id = (r.get("id") or "").lower() or _slug(f"{persona}-{title}")
            data = DraftInput(persona=persona, title=title, body=body,
                              tags=[t.strip() for t in r.get("tags", "").split(",") if t.strip()],
                              visibility=r.get("visibility") or "public")
            try:
                await self.save_draft(p, doc_id, data, source="import")
                if submit:
                    await self.submit(p, doc_id)
                created.append(doc_id)
            except CMSError as e:
                errors.append({"row": i, "id": doc_id, "error": str(e)})
        return {"created": created, "errors": errors}

    async def seed_from_files(self, docs: list[KnowledgeDoc]) -> int:
        """First start only: Git-managed Markdown becomes published v1. Afterwards the CMS is the source of truth."""
        n = 0
        for doc in docs:
            if await self.store.get_doc(doc.id):
                continue
            ts = now()
            v = DocVersion(doc_id=doc.id, version=1, persona=doc.persona, title=doc.title, body=doc.body,
                           tags=doc.tags, visibility=doc.visibility, state="published", source="git", author="git",
                           created_at=ts, updated_at=ts)
            await self.store.insert_version(v)
            await self.store.upsert_doc(DocRecord(id=doc.id, published_version=1))
            n += 1
        return n

    # ---------------------------------------------------------------- helpers
    async def _publish(self, v: DocVersion) -> None:
        for old in await self.store.versions(v.doc_id):
            if old.state == "published" and old.version != v.version:
                old.state = "superseded"
                await self.store.update_version(old)
        v.state, v.updated_at = "published", now()
        await self.store.update_version(v)
        await self.store.upsert_doc(DocRecord(id=v.doc_id, status="active", published_version=v.version))
        await self.on_publish()

    async def _doc(self, doc_id: str) -> DocRecord:
        d = await self.store.get_doc(doc_id)
        if not d:
            raise CMSError("document not found", 404)
        return d

    async def _version(self, doc_id: str, version: int) -> DocVersion:
        v = next((x for x in await self.store.versions(doc_id) if x.version == version), None)
        if not v:
            raise CMSError("version not found", 404)
        return v

    async def _open(self, doc_id: str, states: tuple[str, ...]) -> DocVersion:
        v = next((x for x in await self.store.versions(doc_id) if x.state in states), None)
        if not v:
            raise CMSError(f"no {'/'.join(states)} version", 409)
        return v

    @staticmethod
    def _validate(doc_id: str, d: DraftInput) -> None:
        if not _ID_RE.match(doc_id):
            raise CMSError("id: lowercase letters, digits and '-' (2-80 chars)")
        if d.persona not in PERSONA_VALUES:
            raise CMSError(f"persona must be one of {PERSONA_VALUES}")
        if d.visibility not in ("public", "internal"):
            raise CMSError("visibility must be public or internal")
        if not 3 <= len(d.title) <= 200:
            raise CMSError("title: 3-200 characters")
        if not 10 <= len(d.body) <= 50_000:
            raise CMSError("body: 10-50000 characters")


def parse_upload(filename: str, data: bytes) -> list[dict[str, str]]:
    """Rows from .xlsx (first sheet) or .csv (UTF-8). Google Sheets: File → Download → .xlsx/.csv."""
    name = filename.lower()
    if name.endswith(".csv"):
        return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip() for h in rows[0]]
        return [{h: "" if v is None else str(v) for h, v in zip(header, r)} for r in rows[1:] if any(r)]
    raise CMSError("upload a .xlsx or .csv file")


_TRANSLIT = str.maketrans({
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "yo", "ж": "j", "з": "z", "и": "i",
    "й": "i", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "ө": "o", "п": "p", "р": "r", "с": "s",
    "т": "t", "у": "u", "ү": "u", "ф": "f", "х": "h", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "sh", "ъ": "",
    "ы": "i", "ь": "", "э": "e", "ю": "yu", "я": "ya",
})


def _slug(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", text.lower().translate(_TRANSLIT)).strip("-")
    return s[:60].rstrip("-") or "faq"


class InstructionsService(CMSService):
    """Operation-team bot instructions: one document per persona (id = persona), same
    draft → review → publish → rollback workflow and audit trail as knowledge."""

    audit_prefix = "prompt"
    PERSONAS = ("driver", "client", "customer")

    async def published(self) -> dict[str, str]:
        """persona -> published instruction text."""
        out = {}
        for doc_id, vs in (await self.store.latest_versions()).items():
            pub = next((v for v in vs if v.state == "published"), None)
            if pub:
                out[doc_id] = pub.body
        return out

    async def seed_defaults(self, defaults: dict[str, str]) -> int:
        n = 0
        for persona, text in defaults.items():
            if await self.store.get_doc(persona):
                continue
            ts = now()
            v = DocVersion(doc_id=persona, version=1, persona=persona, title=f"{persona} instructions",
                           body=text.strip(), state="published", source="default", author="default",
                           created_at=ts, updated_at=ts)
            await self.store.insert_version(v)
            await self.store.upsert_doc(DocRecord(id=persona, published_version=1))
            n += 1
        return n

    async def set_archived(self, p: Principal, doc_id: str, archived: bool) -> None:
        raise CMSError("bot instructions cannot be archived; publish a new version instead", 409)

    def _validate(self, doc_id: str, d: DraftInput) -> None:
        if doc_id not in self.PERSONAS or d.persona != doc_id:
            raise CMSError(f"instructions exist per persona: {self.PERSONAS}")
        if not 10 <= len(d.body.strip()) <= INSTRUCTIONS_MAX_CHARS:
            raise CMSError(f"instructions: 10-{INSTRUCTIONS_MAX_CHARS} characters")
